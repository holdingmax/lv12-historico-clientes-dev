"""Salida a Excel de la matriz de clientes. Solo genera el archivo; no lee nada de vuelta."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src.formato import MESES_CORTOS

FORMATO_IMPORTE = "#,##0.00"  # Excel lo muestra con el separador de la configuracion regional
TEXTO_TIPO = {"publica": "Pública", "privado": "Privada"}


def matriz_a_excel(matriz: dict, meses: list[int], duplicados: set[int], anio: int) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes por mes"
    encabezados = (
        ["Cliente", "Tipo"] + [f"{MESES_CORTOS[m - 1]} {anio}" for m in meses]
        + ["Total", "Participación", "Posible duplicado"]
    )
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True)
        celda.fill = PatternFill("solid", fgColor="F0F2F6")
        celda.alignment = Alignment(horizontal="center", wrap_text=True)

    total = matriz["total"]
    for c in matriz["clientes"]:
        ws.append(
            [c["nombre"], TEXTO_TIPO.get(c["tipo"], c["tipo"])]
            + [c["por_mes"].get(m) for m in meses]
            + [c["total"], (c["total"] / total) if total else None,
               "Sí" if c["id"] in duplicados else ""]
        )
    ws.append(
        ["Total", ""] + [matriz["totales_mes"][m] for m in meses] + [total, 1 if total else None, ""]
    )
    for celda in ws[ws.max_row]:
        celda.font = Font(bold=True)

    for fila in ws.iter_rows(min_row=2):
        for celda in fila[2:len(meses) + 3]:  # meses + Total
            celda.number_format = FORMATO_IMPORTE
        fila[len(meses) + 3].number_format = "0.0%"  # Participacion
    ws.column_dimensions["A"].width = 42
    for i in range(2, len(encabezados) + 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 16
    ws.freeze_panes = "B2"

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
